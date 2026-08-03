import os
import uuid
import datetime
from flask import Flask, request, jsonify, send_from_directory
from flask_cors import CORS
import openpyxl
from io import BytesIO

# Import configuration and helpers
from firebase_config import db, bucket, use_firebase, MOCK_STORAGE_DIR
from auth import token_required, hash_password, check_password, generate_token

app = Flask(__name__)
# Enable CORS for frontend development
CORS(app)

# Ensure Mock Storage directory exists if in mock mode
if not use_firebase:
    os.makedirs(MOCK_STORAGE_DIR, exist_ok=True)

# Helper to seed default admin
def seed_admin():
    try:
        admins_ref = db.collection('Admins')
        admins = admins_ref.limit(1).get()
        if not admins:
            admin_id = str(uuid.uuid4())
            admin_data = {
                'adminId': admin_id,
                'email': 'admin@ieee.org',
                'name': 'IEEE Admin',
                'password': hash_password('Admin@123'),
                'role': 'admin'
            }
            admins_ref.document(f'admin_admin@ieee.org').set(admin_data)
            print("Successfully seeded default admin: admin@ieee.org / Admin@123")
    except Exception as e:
        print(f"Error seeding admin: {e}")

seed_admin()

# ----------------------------------------------------
# LOCAL STORAGE SERVING (MOCK MODE FALLBACK)
# ----------------------------------------------------
if not use_firebase:
    @app.route('/api/storage/file/<filename>', methods=['GET'])
    def serve_mock_file(filename):
        return send_from_directory(MOCK_STORAGE_DIR, filename)

# ----------------------------------------------------
# AUTHENTICATION API
# ----------------------------------------------------
@app.route('/api/auth/login', methods=['POST'])
def login():
    data = request.json or {}
    username = data.get('username', '').strip()  # Reg Number for students, Email for Admins/Volunteers
    password = data.get('password', '').strip()
    
    if not username or not password:
        return jsonify({'message': 'Username/Email and Password are required.'}), 400
        
    # 1. Check if the user is an Admin (Username matches email format, checks Admin collection)
    if '@' in username:
        # Admin check
        admin_doc = db.collection('Admins').document(f'admin_{username}').get()
        if admin_doc.exists:
            admin_data = admin_doc.to_dict()
            if check_password(admin_data['password'], password):
                token = generate_token(admin_data['adminId'], 'admin', email=admin_data['email'])
                return jsonify({
                    'token': token,
                    'role': 'admin',
                    'name': admin_data['name'],
                    'email': admin_data['email']
                })
                
        # Volunteer check
        volunteer_doc = db.collection('Volunteers').document(f'volunteer_{username}').get()
        if volunteer_doc.exists:
            vol_data = volunteer_doc.to_dict()
            if check_password(vol_data['password'], password):
                token = generate_token(vol_data['volunteerId'], 'volunteer', email=vol_data['email'])
                return jsonify({
                    'token': token,
                    'role': 'volunteer',
                    'name': vol_data['name'],
                    'email': vol_data['email']
                })
    else:
        # Student check (Username is registration number)
        student_doc = db.collection('Students').document(f'student_{username}').get()
        if student_doc.exists:
            student_data = student_doc.to_dict()
            if check_password(student_data['password'], password):
                token = generate_token(
                    student_data['studentId'], 
                    'student', 
                    registration_number=student_data['registrationNumber']
                )
                return jsonify({
                    'token': token,
                    'role': 'student',
                    'name': student_data['name'],
                    'registrationNumber': student_data['registrationNumber'],
                    'department': student_data['department']
                })
                
    return jsonify({'message': 'Invalid username or password.'}), 401


# ----------------------------------------------------
# ADMIN PORTAL APIs
# ----------------------------------------------------
@app.route('/api/admin/participants/upload', methods=['POST'])
@token_required(allowed_roles=['admin'])
def upload_participants(current_user):
    if 'file' not in request.files:
        return jsonify({'message': 'No file uploaded.'}), 400
        
    file = request.files['file']
    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({'message': 'Invalid file format. Please upload an Excel file.'}), 400
        
    try:
        wb = openpyxl.load_workbook(BytesIO(file.read()))
        sheet = wb.active
        
        # Read header to map columns dynamically
        header = [cell.value for cell in sheet[1]]
        reg_idx, name_idx, dept_idx, email_idx = -1, -1, -1, -1
        
        for idx, val in enumerate(header):
            if not val:
                continue
            val_str = str(val).strip().lower()
            if 'registration' in val_str or 'reg' in val_str:
                reg_idx = idx
            elif 'name' in val_str:
                name_idx = idx
            elif 'department' in val_str or 'dept' in val_str:
                dept_idx = idx
            elif 'email' in val_str:
                email_idx = idx
                
        if reg_idx == -1 or name_idx == -1 or dept_idx == -1:
            return jsonify({'message': 'Required columns ("Registration Number", "Name", "Department") not found in sheet.'}), 400
            
        students_added = []
        # Process rows
        for row_idx in range(2, sheet.max_row + 1):
            row = [sheet.cell(row=row_idx, column=col_idx).value for col_idx in range(1, len(header) + 1)]
            if not row or not row[reg_idx]:
                continue # Skip empty row
                
            reg_num = str(row[reg_idx]).strip()
            name = str(row[name_idx]).strip()
            dept = str(row[dept_idx]).strip()
            email = str(row[email_idx]).strip() if (email_idx != -1 and row[email_idx]) else ""
            
            # Check if student already exists
            student_ref = db.collection('Students').document(f'student_{reg_num}')
            if not student_ref.get().exists:
                student_id = str(uuid.uuid4())
                qr_token = f"qr_{uuid.uuid4().hex}" # generate random unique QR token
                
                student_data = {
                    'studentId': student_id,
                    'registrationNumber': reg_num,
                    'name': name,
                    'department': dept,
                    'email': email,
                    'password': hash_password('Student@123'), # default password
                    'qrToken': qr_token,
                    'role': 'student'
                }
                student_ref.set(student_data)
                students_added.append({
                    'registrationNumber': reg_num,
                    'name': name,
                    'department': dept
                })
                
        return jsonify({
            'message': f'Successfully imported {len(students_added)} participants.',
            'imported': students_added
        })
        
    except Exception as e:
        return jsonify({'message': f'Error parsing Excel file: {str(e)}'}), 500


@app.route('/api/admin/students', methods=['GET'])
@token_required(allowed_roles=['admin'])
def get_students(current_user):
    try:
        docs = db.collection('Students').stream()
        students = []
        for doc in docs:
            data = doc.to_dict()
            # Omit password for safety
            data.pop('password', None)
            students.append(data)
        return jsonify(students)
    except Exception as e:
        return jsonify({'message': str(e)}), 500


@app.route('/api/admin/volunteers', methods=['GET', 'POST'])
@token_required(allowed_roles=['admin'])
def manage_volunteers(current_user):
    if request.method == 'GET':
        try:
            docs = db.collection('Volunteers').stream()
            volunteers = []
            for doc in docs:
                data = doc.to_dict()
                data.pop('password', None)
                volunteers.append(data)
            return jsonify(volunteers)
        except Exception as e:
            return jsonify({'message': str(e)}), 500
            
    elif request.method == 'POST':
        data = request.json or {}
        name = data.get('name', '').strip()
        email = data.get('email', '').strip()
        password = data.get('password', '').strip()
        
        if not name or not email or not password:
            return jsonify({'message': 'All fields (name, email, password) are required.'}), 400
            
        try:
            vol_ref = db.collection('Volunteers').document(f'volunteer_{email}')
            if vol_ref.get().exists:
                return jsonify({'message': 'Volunteer with this email already exists.'}), 400
                
            volunteer_id = str(uuid.uuid4())
            vol_data = {
                'volunteerId': volunteer_id,
                'name': name,
                'email': email,
                'password': hash_password(password),
                'role': 'volunteer'
            }
            vol_ref.set(vol_data)
            return jsonify({'message': f'Volunteer "{name}" created successfully.', 'volunteerId': volunteer_id})
        except Exception as e:
            return jsonify({'message': str(e)}), 500


@app.route('/api/admin/events', methods=['GET', 'POST'])
@token_required(allowed_roles=['admin', 'volunteer'])
def manage_events(current_user):
    if request.method == 'GET':
        try:
            docs = db.collection('Events').stream()
            events = []
            for doc in docs:
                events.append(doc.to_dict())
            return jsonify(events)
        except Exception as e:
            return jsonify({'message': str(e)}), 500
            
    elif request.method == 'POST':
        # Only admins can create events
        if current_user['role'] != 'admin':
            return jsonify({'message': 'Unauthorized action.'}), 403
            
        data = request.json or {}
        event_name = data.get('eventName', '').strip()
        start_date_str = data.get('startDate', '').strip()
        end_date_str = data.get('endDate', '').strip()
        
        if not event_name or not start_date_str or not end_date_str:
            return jsonify({'message': 'Event Name, Start Date, and End Date are required.'}), 400
            
        try:
            start_date = datetime.datetime.strptime(start_date_str, '%Y-%m-%d')
            end_date = datetime.datetime.strptime(end_date_str, '%Y-%m-%d')
            if end_date < start_date:
                return jsonify({'message': 'End Date cannot be before Start Date.'}), 400
                
            delta = end_date - start_date
            days_count = delta.days + 1
            
            event_id = str(uuid.uuid4())
            days_list = [f"Day {i}" for i in range(1, days_count + 1)]
            
            event_data = {
                'eventId': event_id,
                'eventName': event_name,
                'startDate': start_date_str,
                'endDate': end_date_str,
                'daysCount': days_count,
                'days': days_list
            }
            
            db.collection('Events').document(f'event_{event_id}').set(event_data)
            return jsonify({'message': f'Event "{event_name}" created successfully.', 'event': event_data})
        except ValueError:
            return jsonify({'message': 'Invalid date format. Use YYYY-MM-DD.'}), 400
        except Exception as e:
            return jsonify({'message': str(e)}), 500


@app.route('/api/admin/verification/pending', methods=['GET'])
@token_required(allowed_roles=['admin'])
def get_pending_verification(current_user):
    try:
        docs = db.collection('Attendance').where('status', '==', 'Pending Verification').stream()
        records = []
        for doc in docs:
            records.append(doc.to_dict())
        return jsonify(records)
    except Exception as e:
        return jsonify({'message': str(e)}), 500


@app.route('/api/admin/verification/verify', methods=['POST'])
@token_required(allowed_roles=['admin'])
def verify_attendance(current_user):
    data = request.json or {}
    attendance_id = data.get('attendanceId', '').strip()
    status = data.get('status', '').strip() # Verified or Rejected
    remarks = data.get('remarks', '').strip()
    
    if not attendance_id or status not in ['Verified', 'Rejected']:
        return jsonify({'message': 'Attendance ID and valid Status (Verified/Rejected) are required.'}), 400
        
    try:
        att_ref = db.collection('Attendance').document(attendance_id)
        att_doc = att_ref.get()
        if not att_doc.exists:
            return jsonify({'message': 'Attendance record not found.'}), 404
            
        att_ref.update({
            'status': status,
            'remarks': remarks
        })
        return jsonify({'message': f'Attendance record marked as {status}.'})
    except Exception as e:
        return jsonify({'message': str(e)}), 500


@app.route('/api/admin/reports/stats', methods=['GET'])
@token_required(allowed_roles=['admin'])
def get_report_stats(current_user):
    event_id = request.args.get('eventId')
    try:
        # Total registered students
        students = db.collection('Students').get()
        total_participants = len(students)
        
        # Base attendance query
        attendance_query = db.collection('Attendance')
        if event_id:
            attendance_query = attendance_query.where('eventId', '==', event_id)
            
        attendance_records = attendance_query.get()
        
        verified = 0
        pending = 0
        rejected = 0
        day_stats = {}
        
        for doc in attendance_records:
            data = doc.to_dict()
            status = data.get('status')
            day = data.get('day', 'Day 1')
            
            if status == 'Verified':
                verified += 1
            elif status == 'Pending Verification':
                pending += 1
            elif status == 'Rejected':
                rejected += 1
                
            day_stats[day] = day_stats.get(day, 0) + (1 if status == 'Verified' else 0)
            
        return jsonify({
            'totalParticipants': total_participants,
            'verified': verified,
            'pending': pending,
            'rejected': rejected,
            'dayStats': day_stats
        })
    except Exception as e:
        return jsonify({'message': str(e)}), 500


@app.route('/api/admin/reports/export', methods=['GET'])
@token_required(allowed_roles=['admin'])
def export_reports(current_user):
    event_id = request.args.get('eventId')
    if not event_id:
        return jsonify({'message': 'Event ID parameter is required.'}), 400
        
    try:
        event_doc = db.collection('Events').document(f'event_{event_id}').get()
        if not event_doc.exists:
            return jsonify({'message': 'Event not found.'}), 404
        event_name = event_doc.to_dict().get('eventName', 'Event')
        
        # Get students
        students_docs = db.collection('Students').stream()
        students_map = {}
        for s in students_docs:
            sd = s.to_dict()
            students_map[sd['studentId']] = sd
            
        # Get attendance records for this event
        att_docs = db.collection('Attendance').where('eventId', '==', event_id).stream()
        att_map = {} # studentId -> { day: status }
        for a in att_docs:
            ad = a.to_dict()
            s_id = ad['studentId']
            day = ad['day']
            status = ad['status']
            if s_id not in att_map:
                att_map[s_id] = {}
            att_map[s_id][day] = status
            
        # Create Excel workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance Report"
        
        # Write Title
        ws.merge_cells('A1:G1')
        ws['A1'] = f"Attendance Report: {event_name}"
        ws['A1'].font = openpyxl.styles.Font(size=14, bold=True)
        ws['A1'].alignment = openpyxl.styles.Alignment(horizontal='center')
        
        # Write Headers
        headers = ["Registration Number", "Name", "Department", "Email", "Day 1", "Day 2", "Day 3"]
        ws.append([]) # Empty row 2
        ws.append(headers) # Row 3
        
        # Style Header Row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=3, column=col)
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill(start_color="00629B", end_color="00629B", fill_type="solid")
            
        # Write student attendance data
        for s_id, s_data in students_map.items():
            s_att = att_map.get(s_id, {})
            row = [
                s_data.get('registrationNumber'),
                s_data.get('name'),
                s_data.get('department'),
                s_data.get('email', ''),
                s_att.get('Day 1', 'Absent'),
                s_att.get('Day 2', 'Absent'),
                s_att.get('Day 3', 'Absent')
            ]
            ws.append(row)
            
        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # Return Excel file
        from flask import make_response
        response = make_response(output.read())
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        response.headers['Content-Disposition'] = f'attachment; filename=IEEE_Report_{event_id}.xlsx'
        return response
        
    except Exception as e:
        return jsonify({'message': f'Error exporting report: {str(e)}'}), 500


@app.route('/api/admin/attendance', methods=['GET'])
@token_required(allowed_roles=['admin'])
def get_attendance_records(current_user):
    event_id = request.args.get('eventId')
    if not event_id:
        return jsonify({'message': 'Event ID is required.'}), 400
    try:
        docs = db.collection('Attendance').where('eventId', '==', event_id).stream()
        records = []
        for doc in docs:
            records.append(doc.to_dict())
        return jsonify(records)
    except Exception as e:
        return jsonify({'message': str(e)}), 500


# ----------------------------------------------------
# STUDENT PORTAL APIs
# ----------------------------------------------------
@app.route('/api/student/profile', methods=['GET'])
@token_required(allowed_roles=['student'])
def get_student_profile(current_user):
    reg_num = current_user['registrationNumber']
    try:
        student_doc = db.collection('Students').document(f'student_{reg_num}').get()
        if not student_doc.exists:
            return jsonify({'message': 'Student record not found.'}), 404
            
        profile = student_doc.to_dict()
        profile.pop('password', None)
        
        # Get active events and student attendance records
        events_docs = db.collection('Events').stream()
        events = [e.to_dict() for e in events_docs]
        
        attendance_docs = db.collection('Attendance').where('studentId', '==', profile['studentId']).stream()
        attendance_records = [a.to_dict() for a in attendance_docs]
        
        # Build checklist of attendance per event-day
        events_checklist = []
        for ev in events:
            ev_id = ev['eventId']
            ev_name = ev['eventName']
            
            day_statuses = {}
            for day in ev.get('days', ['Day 1', 'Day 2', 'Day 3']):
                # Find matching attendance record
                status = 'Absent'
                for att in attendance_records:
                    if att['eventId'] == ev_id and att['day'] == day:
                        status = att['status']
                        break
                day_statuses[day] = status
                
            events_checklist.append({
                'eventId': ev_id,
                'eventName': ev_name,
                'attendance': day_statuses
            })
            
        profile['attendanceChecklist'] = events_checklist
        return jsonify(profile)
        
    except Exception as e:
        return jsonify({'message': str(e)}), 500


# ----------------------------------------------------
# VOLUNTEER PORTAL APIs
# ----------------------------------------------------
@app.route('/api/volunteer/scan', methods=['POST'])
@token_required(allowed_roles=['volunteer'])
def scan_qr_token(current_user):
    data = request.json or {}
    qr_token = data.get('qrToken', '').strip()
    
    if not qr_token:
        return jsonify({'message': 'QR Token is required.'}), 400
        
    try:
        docs = db.collection('Students').where('qrToken', '==', qr_token).limit(1).get()
        if not docs:
            return jsonify({'message': 'Invalid QR Token. Student not found.'}), 404
            
        student_data = docs[0].to_dict()
        student_data.pop('password', None)
        return jsonify(student_data)
        
    except Exception as e:
        return jsonify({'message': str(e)}), 500


@app.route('/api/volunteer/submit-attendance', methods=['POST'])
@token_required(allowed_roles=['volunteer'])
def submit_attendance(current_user):
    data = request.json or {}
    student_id = data.get('studentId', '').strip()
    event_id = data.get('eventId', '').strip()
    day = data.get('day', '').strip()
    face_photo_base64 = data.get('facePhoto', '') # Base64 encoded JPEG
    id_photo_base64 = data.get('idPhoto', '') # Base64 encoded JPEG
    
    if not student_id or not event_id or not day or not face_photo_base64 or not id_photo_base64:
        return jsonify({'message': 'Required fields missing: studentId, eventId, day, facePhoto, idPhoto.'}), 400
        
    try:
        # Check if student exists
        student_query = db.collection('Students').where('studentId', '==', student_id).limit(1).get()
        if not student_query:
            return jsonify({'message': 'Student not found.'}), 404
        student_data = student_query[0].to_dict()
        
        # Check if event exists
        event_doc = db.collection('Events').document(f'event_{event_id}').get()
        if not event_doc.exists:
            return jsonify({'message': 'Event not found.'}), 404
        event_name = event_doc.to_dict().get('eventName')
        
        # Check duplicate attendance using unique document ID
        att_id = f"attendance_{event_id}_{day}_{student_id}"
        att_ref = db.collection('Attendance').document(att_id)
        if att_ref.get().exists:
            return jsonify({'message': 'Attendance already recorded for this student on this day.'}), 400
            
        # 1. Upload facePhoto to Firebase Storage
        face_blob_path = f"attendance/{event_id}/{day}/{student_id}_face.jpg"
        face_blob = bucket.blob(face_blob_path)
        face_blob.upload_from_string(face_photo_base64, content_type='image/jpeg')
        face_blob.make_public()
        face_photo_url = face_blob.public_url
        
        # 2. Upload idPhoto to Firebase Storage
        id_blob_path = f"attendance/{event_id}/{day}/{student_id}_id.jpg"
        id_blob = bucket.blob(id_blob_path)
        id_blob.upload_from_string(id_photo_base64, content_type='image/jpeg')
        id_blob.make_public()
        id_photo_url = id_blob.public_url
        
        # Save attendance record
        attendance_record = {
            'attendanceId': att_id,
            'studentId': student_id,
            'registrationNumber': student_data['registrationNumber'],
            'name': student_data['name'],
            'department': student_data['department'],
            'eventId': event_id,
            'eventName': event_name,
            'day': day,
            'facePhotoURL': face_photo_url,
            'idPhotoURL': id_photo_url,
            'status': 'Pending Verification',
            'remarks': '',
            'volunteerId': current_user['userId'],
            'timestamp': datetime.datetime.utcnow().isoformat()
        }
        
        att_ref.set(attendance_record)
        return jsonify({'message': 'Attendance submitted successfully.', 'record': attendance_record})
        
    except Exception as e:
        return jsonify({'message': f'Server error submitting attendance: {str(e)}'}), 500


@app.route('/api/auth/change-password', methods=['POST'])
@token_required
def change_password(current_user):
    data = request.json or {}
    old_password = data.get('oldPassword', '').strip()
    new_password = data.get('newPassword', '').strip()
    
    if not old_password or not new_password:
        return jsonify({'message': 'Old and New passwords are required.'}), 400
        
    try:
        # Check database based on role
        if current_user['role'] == 'admin':
            user_ref = db.collection('Admins').document(f"admin_{current_user['email']}")
        elif current_user['role'] == 'volunteer':
            user_ref = db.collection('Volunteers').document(f"volunteer_{current_user['email']}")
        elif current_user['role'] == 'student':
            user_ref = db.collection('Students').document(f"student_{current_user['registrationNumber']}")
        else:
            return jsonify({'message': 'Unauthorized role.'}), 403
            
        user_doc = user_ref.get()
        if not user_doc.exists:
            return jsonify({'message': 'User record not found.'}), 404
            
        user_data = user_doc.to_dict()
        if not check_password(user_data['password'], old_password):
            return jsonify({'message': 'Incorrect old password.'}), 400
            
        user_ref.update({
            'password': hash_password(new_password)
        })
        return jsonify({'message': 'Password updated successfully!'})
        
    except Exception as e:
        return jsonify({'message': str(e)}), 500


if __name__ == "__main__":
    app.run(debug=True, port=5000)